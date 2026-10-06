#!/usr/bin/env python3
# classify_assets_iso.py — Clasificación normativa ISO 27001 para inventario de activos
import openpyxl
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from _paths import EXCEL, EXPORTS, SNAPSHOTS  # noqa: F401
SHEET_MASTER = "09_analysis_master"
SHEET_ISO = "12_inventario_iso"

# Estilos de celdas
FILL_HEADER = PatternFill("solid", fgColor="1F4E79")
FILL_RED = PatternFill("solid", fgColor="FADBD8")    # Confidencial / SI
FILL_YELLOW = PatternFill("solid", fgColor="FCF3CF") # Restringido
FILL_GREEN = PatternFill("solid", fgColor="D5F5E3")  # Uso Interno / NO

FONT_WHITE_BOLD = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
FONT_REGULAR = Font(name="Segoe UI", size=10)
FONT_BOLD = Font(name="Segoe UI", size=10, bold=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3")
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

def classify_row(row_data, col_idx):
    r_type = (row_data[col_idx["type"]-1] or "").lower()
    env = (row_data[col_idx["environment"]-1] or "").lower()
    sub_name = (row_data[col_idx["subscription_name"]-1] or "").lower()
    owner = row_data[col_idx["owner_confirmed"]-1] or "sin-custodio"
    
    # Determinar si es producción o QA de producción
    is_prod = "prod" in env or "production" in env or "qa-produccion" in sub_name
    
    # Identificar recursos de base de datos / llaves
    is_data_key = any(t in r_type for t in ["sql", "vault", "storage", "cosmos", "documentdb"])
    is_compute_net = any(t in r_type for t in ["containerservice", "virtualmachines", "web/sites", "networksecuritygroups"])

    if is_prod:
        if is_data_key:
            return "Confidencial", owner, 3, 3, 3, 9, "SI"
        elif is_compute_net:
            return "Restringido", owner, 2, 3, 3, 8, "SI"
        else:
            return "Restringido", owner, 2, 2, 2, 6, "SI"
    else:
        # Dev / QA / Test
        if is_data_key:
            return "Restringido", owner, 2, 2, 2, 6, "NO"
        else:
            return "Uso Interno", owner, 1, 1, 1, 3, "NO"

def main():
    if not EXCEL.exists():
        print(f"Error: No se encontró el Excel en {EXCEL}")
        return

    wb = openpyxl.load_workbook(EXCEL)
    if SHEET_MASTER not in wb.sheetnames:
        print(f"Error: No existe la hoja {SHEET_MASTER} en {EXCEL}")
        return

    ws_master = wb[SHEET_MASTER]
    headers = [c.value for c in ws_master[1]]
    col_idx = {h: i+1 for i, h in enumerate(headers) if h}

    # Eliminar hoja ISO si ya existe
    if SHEET_ISO in wb.sheetnames:
        del wb[SHEET_ISO]
    
    ws_iso = wb.create_sheet(SHEET_ISO)
    ws_iso.views.sheetView[0].showGridLines = True

    # Cabeceras del nuevo reporte ISO 27001
    iso_headers = [
        "name", "resourceType", "location", "resourceGroup", "subscription", "Idtags",
        "Clasificación del activo", "custodio", "confidencialidad", "disponibilidad",
        "integridad", "puntuación del activo", "Gestión de riesgo (SI/NO)"
    ]

    # Escribir cabecera
    for col_num, header in enumerate(iso_headers, 1):
        cell = ws_iso.cell(row=1, column=col_num, value=header)
        cell.font = FONT_WHITE_BOLD
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
        
    # Clasificar recursos
    row_count = 1
    for r_idx in range(2, ws_master.max_row + 1):
        row_cells = [ws_master.cell(row=r_idx, column=c).value for c in range(1, ws_master.max_column + 1)]
        # Saltar si no hay nombre o tipo
        if not row_cells[col_idx["name"]-1] or not row_cells[col_idx["type"]-1]:
            continue
            
        name = row_cells[col_idx["name"]-1]
        r_type = row_cells[col_idx["type"]-1]
        location = row_cells[col_idx["location"]-1]
        rg = row_cells[col_idx["resourceGroup"]-1]
        sub = row_cells[col_idx["subscription_name"]-1]
        tags = row_cells[col_idx["tags"]-1] or ""

        classification, custodio, c_val, d_val, i_val, score, risk_mgmt = classify_row(row_cells, col_idx)
        
        row_count += 1
        # Escribir valores
        ws_iso.cell(row=row_count, column=1, value=name).font = FONT_BOLD
        ws_iso.cell(row=row_count, column=2, value=r_type).font = FONT_REGULAR
        ws_iso.cell(row=row_count, column=3, value=location).font = FONT_REGULAR
        ws_iso.cell(row=row_count, column=4, value=rg).font = FONT_REGULAR
        ws_iso.cell(row=row_count, column=5, value=sub).font = FONT_REGULAR
        ws_iso.cell(row=row_count, column=6, value=tags).font = FONT_REGULAR
        
        c_class = ws_iso.cell(row=row_count, column=7, value=classification)
        c_class.font = FONT_BOLD
        c_class.alignment = ALIGN_CENTER
        
        ws_iso.cell(row=row_count, column=8, value=custodio).font = FONT_REGULAR
        
        c_val_cell = ws_iso.cell(row=row_count, column=9, value=c_val)
        c_val_cell.font = FONT_REGULAR
        c_val_cell.alignment = ALIGN_CENTER
        
        d_val_cell = ws_iso.cell(row=row_count, column=10, value=d_val)
        d_val_cell.font = FONT_REGULAR
        d_val_cell.alignment = ALIGN_CENTER
        
        i_val_cell = ws_iso.cell(row=row_count, column=11, value=i_val)
        i_val_cell.font = FONT_REGULAR
        i_val_cell.alignment = ALIGN_CENTER
        
        score_cell = ws_iso.cell(row=row_count, column=12, value=score)
        score_cell.font = FONT_BOLD
        score_cell.alignment = ALIGN_CENTER
        
        risk_cell = ws_iso.cell(row=row_count, column=13, value=risk_mgmt)
        risk_cell.font = FONT_BOLD
        risk_cell.alignment = ALIGN_CENTER

        # Aplicar bordes a toda la fila
        for col_num in range(1, 14):
            ws_iso.cell(row=row_count, column=col_num).border = THIN_BORDER
            
        # Colorear según clasificación
        if classification == "Confidencial":
            c_class.fill = FILL_RED
            score_cell.fill = FILL_RED
        elif classification == "Restringido":
            c_class.fill = FILL_YELLOW
            score_cell.fill = FILL_YELLOW
        else:
            c_class.fill = FILL_GREEN
            score_cell.fill = FILL_GREEN

        # Colorear Gestión de Riesgo
        if risk_mgmt == "SI":
            risk_cell.fill = FILL_RED
        else:
            risk_cell.fill = FILL_GREEN

    # Auto-ajustar ancho de columnas
    for col in ws_iso.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws_iso.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

    wb.save(EXCEL)
    print(f"✅ Hoja '{SHEET_ISO}' creada exitosamente con {row_count-1} activos clasificados.")

if __name__ == "__main__":
    main()
