#!/usr/bin/env python3
# bootstrap_workbook.py — Crea el Excel maestro vacio si todavia no existe.
#
# El resto del pipeline abre el libro con `load_workbook` y espera encontrar la
# hoja 09_analysis_master con sus columnas. En una instalacion nueva ese libro
# no existe: este paso lo crea con la estructura minima, y no toca nada si ya
# hay uno (el master acumula revisiones manuales que no deben perderse).
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from _paths import EXCEL

SHEET_MASTER = "09_analysis_master"

COLUMNAS = [
    # Identidad del recurso (las sincroniza sync_master desde Azure)
    "name", "type", "resourceGroup", "location", "subscriptionId",
    "subscription_name", "resource_id", "tags",
    # Clasificacion (enrich_master, normalize_classify, provisioning_classify)
    "domain", "subdomain", "environment", "provisioning_method",
    "evidence_source", "known_drift",
    # Custodia y adopcion de IaC (tags_enrich, fix_owner_pm, finalize_master)
    "owner_confirmed", "candidate_module", "adoption_priority",
    "requires_terraform_import",
    # Trazabilidad
    "extraction_date", "inventory_version", "review_status",
]


def main():
    if EXCEL.exists():
        print(f"  = {EXCEL.name} ya existe; no se modifica.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_MASTER
    relleno = PatternFill("solid", fgColor="1F4E78")
    for i, columna in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=i, value=columna)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    EXCEL.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL)
    print(f"  + Creado {EXCEL} con la hoja {SHEET_MASTER}.")


if __name__ == "__main__":
    main()
