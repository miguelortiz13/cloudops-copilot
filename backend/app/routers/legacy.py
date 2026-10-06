
from fastapi import APIRouter, HTTPException
from typing import Optional
from app.schemas.inventory import *
from app.schemas.k8s import *
from app.schemas.requests import *
from app.core import config
from app.core.container import get_services, sync_status
from app.routers.teams import get_microsoft_oauth_token
import os
import requests


router = APIRouter(tags=['legacy'])

@router.get("/")
def read_root():
    return {
        "status": "online", 
        "agent": "Azure Inventory AI Bot Service",
        "mode": "Azure Live Cloud Query" if get_services()[0].azure_connected else "Unauthenticated"
    }


@router.get("/api/stats")
def get_stats(subscriptions: Optional[str] = None):
    try:
        subs_list = subscriptions.split(",") if subscriptions else None
        stats = get_services()[0].get_summary_stats(subscriptions=subs_list)
        return stats
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/api/resources")
def get_resources(
    query: Optional[str] = None, 
    provisioning: Optional[str] = None,
    location: Optional[str] = None,
    tagging_ok: Optional[str] = None,
    subscriptions: Optional[str] = None,
    page: Optional[int] = None,
    page_size: int = 50,
    limit: int = 2000
):
    try:
        subs_list = subscriptions.split(",") if subscriptions else None
        search_limit = 5000 if page is not None else limit
        results = get_services()[0].search_resources(query or "", limit=search_limit, subscriptions=subs_list)
        
        if provisioning:
            results = [r for r in results if r.get('provisioning_method', '').lower() == provisioning.lower()]
            
        if location:
            results = [r for r in results if r.get('location', '').lower() == location.lower()]
            
        if tagging_ok:
            results = [r for r in results if r.get('tagging_ok', '').lower() == tagging_ok.lower()]
            
        total = len(results)
        
        if page is not None:
            offset = (page - 1) * page_size
            paginated_results = results[offset : offset + page_size]
            return {
                "data": paginated_results,
                "total": total,
                "page": page,
                "page_size": page_size
            }
            
        return results
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/api/governance/iso")
def get_iso_governance_report():
    import openpyxl
    
    excel_path = config.INVENTORY_EXCEL
    if not excel_path.exists():
        return {"error": "El archivo de inventario no ha sido generado todavía.", "data": [], "stats": {}}
        
    try:
        # Load in read-only to be fast and safe
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        if "12_inventario_iso" not in wb.sheetnames:
            return {"error": "La clasificación ISO no ha sido generada todavía.", "data": [], "stats": {}}
            
        ws = wb["12_inventario_iso"]
        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            return {"error": "La hoja ISO está vacía.", "data": [], "stats": {}}
            
        headers = rows[0]
        data = []
        for r in rows[1:]:
            if not r[0]:  # Skip if name is empty
                continue
            data.append({headers[i]: r[i] for i in range(len(headers))})
            
        # Calculate summaries for statistics
        classification_counts = {}
        risk_counts = {"SI": 0, "NO": 0}
        score_sum = 0
        total_assets = len(data)
        
        for item in data:
            c = item.get("Clasificación del activo") or "Uso Interno"
            classification_counts[c] = classification_counts.get(c, 0) + 1
            
            r = item.get("Gestión de riesgo (SI/NO)") or "NO"
            risk_counts[r] = risk_counts.get(r, 0) + 1
            
            score_sum += int(item.get("puntuación del activo") or 3)
            
        avg_score = round(score_sum / total_assets, 2) if total_assets > 0 else 0
        
        return {
            "data": data,
            "stats": {
                "total_assets": total_assets,
                "classification_distribution": classification_counts,
                "risk_management_distribution": risk_counts,
                "average_criticality_score": avg_score
            }
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/api/subscriptions")
def get_subscriptions():
    """Queries Azure Resource Graph to list all subscriptions accessible to the Service Principal."""
    try:
        kql = "resourcecontainers | where type == 'microsoft.resources/subscriptions' | project name, subscriptionId"
        raw = get_services()[0].query_azure_resource_graph(kql, bypass_cache=False, subscriptions=[])
        return [{"name": r.get("name"), "id": r.get("subscriptionId")} for r in raw]
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

def send_pipeline_report_to_teams(webhook_url: str, n_added: int, n_removed: int):
    from datetime import date
    
    try:
        stats = get_services()[0].get_summary_stats()
        total_resources = stats.get("total_resources", 0)
        tag_compliance = stats.get("tag_compliance_percentage", 0.0)
    except Exception:
        total_resources = "N/A"
        tag_compliance = "N/A"
        
    title = "📊 Reporte de Sincronización: Inventario Semanal Azure"
    color = "10B981" if n_added == 0 and n_removed == 0 else "3B82F6"
    
    facts = [
        {"name": "Fecha de Proceso", "value": str(date.today())},
        {"name": "Recursos Nuevos Detectados", "value": f"🟢 +{n_added} recursos"},
        {"name": "Recursos Eliminados (Bajas)", "value": f"🔴 -{n_removed} recursos"},
        {"name": "Total Activos en Azure", "value": f"{total_resources} recursos"},
        {"name": "Cumplimiento de Tags", "value": f"{tag_compliance}%"}
    ]
    
    payload = {
        "@type": "MessageCard",
        "@context": "http://schema.org/extensions",
        "themeColor": color,
        "summary": title,
        "title": title,
        "sections": [{
            "activityTitle": "DevOps Agents — Pipeline de Visibilidad Cloud",
            "activitySubtitle": "Actualización exitosa del Excel Maestro e Histórico de Snapshots",
            "activityImage": "https://img.icons8.com/color/96/azure-infrastructure.png",
            "facts": facts,
            "markdown": True
        }],
        "potentialAction": [{
            "@type": "OpenUri",
            "name": "Ver Panel de Control",
            "targets": [{"os": "default", "uri": config.FRONTEND_URL}]
        }]
    }
    
    try:
        requests.post(webhook_url, json=payload, headers={"Content-Type": "application/json"}, timeout=12)
    except Exception as err:
        print(f"Failed to post weekly inventory notification to Teams: {err}")

def send_proactive_report_to_teams(service_url: str, conversation_id: str, n_added: int, n_removed: int):
    from datetime import date
    
    token = get_microsoft_oauth_token()
    if not token:
        print("❌ Cannot send proactive report: Failed to acquire Microsoft OAuth token.")
        return
        
    try:
        stats = get_services()[0].get_summary_stats()
        total_resources = stats.get("total_resources", 0)
        tag_compliance = stats.get("tag_compliance_percentage", 0.0)
    except Exception:
        total_resources = "N/A"
        tag_compliance = "N/A"
        
    title = "📊 Reporte de Sincronización: Inventario Semanal Azure"
    
    card_text = (
        f"**Fecha de Proceso:** {date.today()}\n\n"
        f"🟢 **Recursos Nuevos Detectados:** +{n_added} recursos\n"
        f"🔴 **Recursos Eliminados (Bajas):** -{n_removed} recursos\n\n"
        f"**Total Activos en Azure:** {total_resources} recursos\n"
        f"**Cumplimiento de Tags:** {tag_compliance}%"
    )
    
    payload = {
        "type": "message",
        "attachments": [
            {
                "contentType": "application/vnd.microsoft.card.hero",
                "content": {
                    "title": title,
                    "subtitle": "Actualización exitosa del Excel Maestro e Histórico de Snapshots",
                    "text": card_text,
                    "buttons": [
                        {
                            "type": "openUrl",
                            "title": "Ver Panel de Control",
                            "value": config.FRONTEND_URL
                        }
                    ]
                }
            }
        ]
    }
    
    url = f"{service_url.rstrip('/')}/v3/conversations/{conversation_id}/activities"
    headers = {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json"
    }
    
    try:
        res = requests.post(url, json=payload, headers=headers, timeout=12)
        if res.status_code not in [200, 201, 202]:
            print(f"❌ Failed to post proactive message. Status: {res.status_code} - {res.text}")
        else:
            print("📬 Proactive report successfully sent via Teams Bot.")
    except Exception as err:
        print(f"❌ Exception posting proactive Teams Bot message: {err}")

def run_inventory_pipeline():
    import subprocess
    import re
    import json
    from datetime import date
    
    sync_status["running"] = True
    sync_status["logs"] = "Iniciando pipeline de inventario semanal en la nube...\n"
    sync_status["error"] = None
    
    pipeline_dir = config.PIPELINE_DIR
    log_dir = config.SNAPSHOTS_DIR
    log_dir.mkdir(parents=True, exist_ok=True)
    
    today_str = date.today().isoformat()
    log_file = log_dir / f"weekly_{today_str}.log"
    sync_status["log_file"] = str(log_file)
    
    try:
        with open(log_file, "w") as f:
            process = subprocess.Popen(
                ["bash", "weekly_inventory.sh"],
                cwd=str(pipeline_dir),
                env={**os.environ, "DATA_DIR": str(config.DATA_DIR)},
                stdout=f,
                stderr=subprocess.STDOUT
            )
            process.wait()
            
        if process.returncode == 0:
            sync_status["running"] = False
            sync_status["last_run"] = today_str
            with open(log_file, "r") as f:
                logs_content = f.read()
                sync_status["logs"] = logs_content
            
            n_added = 0
            n_removed = 0
            try:
                add_match = re.search(r'Nuevos:\s+(\d+)', logs_content)
                if add_match:
                    n_added = int(add_match.group(1))
                    
                del_match = re.search(r'Eliminados:\s+(\d+)', logs_content)
                if del_match:
                    n_removed = int(del_match.group(1))
            except Exception as ex:
                print(f"Error parsing pipeline logs for stats: {ex}")
            
            # Send proactive Teams Bot report if teams_session.json exists
            session_file = config.DATA_DIR / "teams_session.json"
            if session_file.exists():
                try:
                    with open(session_file, "r") as sf:
                        session_data = json.load(sf)
                    
                    s_url = session_data.get("serviceUrl")
                    c_id = session_data.get("conversationId")
                    
                    if s_url and c_id:
                        send_proactive_report_to_teams(s_url, c_id, n_added, n_removed)
                except Exception as ex:
                    print(f"Error sending proactive Teams Bot report: {ex}")
            
            # Send automatic Teams webhook report if TEAMS_WEBHOOK_URL env is set (fallback)
            webhook_url = os.getenv("TEAMS_WEBHOOK_URL")
            if webhook_url:
                try:
                    send_pipeline_report_to_teams(webhook_url, n_added, n_removed)
                except Exception as ex:
                    print(f"Error sending Teams webhook report: {ex}")
        else:
            sync_status["running"] = False
            sync_status["error"] = "Error en el script de orquestación."
            with open(log_file, "r") as f:
                sync_status["logs"] = f.read()
    except Exception as e:
        sync_status["running"] = False
        sync_status["error"] = str(e)
        sync_status["logs"] += f"\n❌ Excepción al ejecutar pipeline: {e}\n"


